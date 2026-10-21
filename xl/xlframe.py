"""Framework: build a workbook of worked examples with live Excel formulas."""
import re
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

FONT = "Arial"
BLUE = Font(name=FONT, color="0000FF")
BLACK = Font(name=FONT, color="000000")
GREEN = Font(name=FONT, color="008000")
BOLD = Font(name=FONT, bold=True)
WHITE_B = Font(name=FONT, bold=True, color="FFFFFF")
ITAL = Font(name=FONT, italic=True, color="444444")
NAVY = PatternFill("solid", fgColor="1F3864")
ORANGE = PatternFill("solid", fgColor="F4B183")
LIGHT = PatternFill("solid", fgColor="FBE5D6")
GREY = PatternFill("solid", fgColor="EDEDED")
INFILL = PatternFill("solid", fgColor="FFF2CC")
thin = Side(style="thin", color="BFBFBF")

FT_CONV = {"m": ("/0.3048", "ft"), "cm": ("/30.48", "ft"), "mm": ("/304.8", "ft"),
           "m/s": ("/0.3048", "ft/s"), "m²": ("/0.09290304", "ft²"), "m³": ("/0.0283168", "ft³"),
           "pts/m²": ("*0.09290304", "pts/ft²")}


class Book:
    def __init__(self):
        self.sheets = []          # list of dict(name, title, examples)

    def sheet(self, name, title, module):
        self.cur_sheet = dict(name=name, title=title, module=module, examples=[])
        self.sheets.append(self.cur_sheet)

    def ex(self, title, given=""):
        s = self.cur_sheet
        num = f"{s['module']}.{len(s['examples']) + 1}" if s['module'] else str(len(s['examples']) + 1)
        self.cur = dict(title=title, given=given, rows=[], num=num, idx=len(s['examples']) + 1)
        s['examples'].append(self.cur)

    def I(self, name, label, value, unit="", note=""):
        self.cur['rows'].append(dict(kind="in", name=name, label=label, value=value, unit=unit, note=note))

    def C(self, name, label, formula, unit="", eq="", book=None, tol=None, fmt=None, note=""):
        self.cur['rows'].append(dict(kind="calc", name=name, label=label, formula=formula, unit=unit,
                                     eq=eq, book=book, tol=tol, fmt=fmt, note=note))

    def T(self, text):
        self.cur['rows'].append(dict(kind="text", label=text))

    # ------------------------------------------------------------------
    def save(self, path, readme_fn=None, extra_fn=None):
        wb = Workbook()
        wb.remove(wb.active)
        index_rows = []
        if readme_fn:
            readme = wb.create_sheet("README")
        for s in self.sheets:
            ws = wb.create_sheet(s['name'])
            self._write_sheet(ws, s, index_rows)
        if extra_fn:
            extra_fn(wb, index_rows)
        if readme_fn:
            readme_fn(readme, index_rows)
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for c in row:
                    if c.font and c.font.name != FONT:
                        f = c.font
                        c.font = Font(name=FONT, bold=f.bold, italic=f.italic, color=f.color, size=f.size, underline=f.underline)
        wb.save(path)

    def _write_sheet(self, ws, s, index_rows):
        widths = [44, 16, 10, 46, 13, 9, 16, 9, 40]
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws["A1"] = s['title']; ws["A1"].font = Font(name=FONT, bold=True, size=14, color="1F3864")
        ws["A2"] = ("Blue = input (edit it)   Black = formula   Green = linked from another example/sheet   "
                    "Check = formula result vs. the value printed in the seminar book")
        ws["A2"].font = ITAL
        hdr = ["Step / quantity", "Value", "Unit", "Equation (as in the book)", "Book value",
               "Check vs. book", "US feet", "Tolerance", "Notes"]
        for j, h in enumerate(hdr, 1):
            c = ws.cell(row=4, column=j, value=h); c.font = WHITE_B; c.fill = NAVY
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.freeze_panes = "A5"
        r = 5
        # pass 1: assign rows
        names = {}
        layout = []
        for e in s['examples']:
            r += 1
            e['hdr_row'] = r; r += 1
            if e['given']:
                e['given_row'] = r; r += 1
            for row in e['rows']:
                row['r'] = r
                if row.get('name'):
                    names[(e['idx'], row['name'])] = r
                r += 1
            e['last_row'] = r - 1
        self._names = names

        def resolve(formula, e_idx):
            def rep(m):
                tok = m.group(1)
                if ":" in tok:
                    a, b = tok.split(":")
                    return f"$B${resolve_one(a, e_idx)}:$B${resolve_one(b, e_idx)}"
                return f"$B${resolve_one(tok, e_idx)}"
            return re.sub(r"\{([A-Za-z0-9_.:]+)\}", rep, formula)

        def resolve_one(tok, e_idx):
            if tok.startswith("X") and "." in tok:
                k, nm = tok[1:].split(".", 1)
                return names[(int(k), nm)]
            return names[(e_idx, tok)]

        for e in s['examples']:
            hr = e['hdr_row']
            ws.cell(row=hr, column=1, value=f"Worked example {e['num']} — {e['title']}").font = WHITE_B
            for j in range(1, 10):
                ws.cell(row=hr, column=j).fill = PatternFill("solid", fgColor="C55A11")
            if e['given']:
                gr = e['given_row']
                ws.cell(row=gr, column=1, value="Given: " + e['given']).font = ITAL
                ws.merge_cells(start_row=gr, start_column=1, end_row=gr, end_column=9)
                ws.cell(row=gr, column=1).alignment = Alignment(wrap_text=True, vertical="top")
                ws.row_dimensions[gr].height = max(15, 15 * (1 + len(e['given']) // 150))
            checks = []
            for row in e['rows']:
                rr = row['r']
                if row['kind'] == "text":
                    c = ws.cell(row=rr, column=1, value=row['label']); c.font = Font(name=FONT, bold=True, color="833C0B")
                    ws.merge_cells(start_row=rr, start_column=1, end_row=rr, end_column=9)
                    c.alignment = Alignment(wrap_text=True, vertical="top")
                    c.fill = LIGHT
                    ws.row_dimensions[rr].height = max(15, 15 * (1 + len(row['label']) // 150))
                    continue
                ws.cell(row=rr, column=1, value=row['label']).font = BLACK
                ws.cell(row=rr, column=1).alignment = Alignment(indent=1, wrap_text=True)
                ws.cell(row=rr, column=3, value=row['unit']).font = BLACK
                if row['kind'] == "in":
                    c = ws.cell(row=rr, column=2, value=row['value']); c.font = BLUE; c.fill = INFILL
                    c.number_format = "General"
                    ws.cell(row=rr, column=4, value="input").font = ITAL
                else:
                    f = "=" + resolve(row['formula'], e['idx'])
                    c = ws.cell(row=rr, column=2, value=f)
                    c.font = GREEN if ("!" in row['formula'] or "{X" in row['formula']) else BLACK
                    c.number_format = row['fmt'] or "0.0000"
                    ws.cell(row=rr, column=4, value=('→ ' + row['eq'][1:].strip()) if row['eq'].startswith('=') else row['eq']).font = BLACK
                    if row['book'] is not None:
                        b = ws.cell(row=rr, column=5, value=row['book']); b.font = BLACK
                        b.number_format = "General"
                        tol = row['tol'] if row['tol'] is not None else 0.0005
                        t = ws.cell(row=rr, column=8, value=tol); t.font = BLUE; t.number_format = "General"
                        ws.cell(row=rr, column=6, value=f'=IF(ABS(B{rr}-E{rr})<=H{rr},"OK","CHECK")').font = BOLD
                        ws.cell(row=rr, column=6).alignment = Alignment(horizontal="center")
                        checks.append(rr)
                u = row['unit']
                if u in FT_CONV and row['kind'] == "calc" or (u in FT_CONV and row['kind'] == "in"):
                    op, fu = FT_CONV[u]
                    g = ws.cell(row=rr, column=7, value=f"=B{rr}{op}")
                    g.number_format = f'0.0000" {fu}"'; g.font = BLACK
                if row.get('note'):
                    ws.cell(row=rr, column=9, value=row['note']).font = ITAL
                ws.cell(row=rr, column=4).alignment = Alignment(wrap_text=True)
                ws.cell(row=rr, column=9).alignment = Alignment(wrap_text=True)
                for j in range(1, 10):
                    ws.cell(row=rr, column=j).border = Border(bottom=thin)
            e['checks'] = checks
            index_rows.append(dict(sheet=s['name'], num=e['num'], title=e['title'], row=hr, last=e['last_row'], checks=checks))
